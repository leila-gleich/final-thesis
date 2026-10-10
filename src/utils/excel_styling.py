"""
src/utils/excel_styling.py
--------------------------
Standardized OpenPyXL Styling and Layout Engine for APA 7th Edition Workbooks.

Enforces:
1. APA 7th Edition Table Formatting Rules:
   - Header row: thin black rule above (top) and below (bottom).
   - Data rows: zero internal horizontal or vertical gridlines.
   - Closing row: thin black rule below (bottom).
   - Zero background fills or zebra striping (clean white/transparent).
2. Consistent Typography (Calibri):
   - Table ID bold on line 1, Table Title italic on line 2.
   - Header row bold, note block below table with italic 'Note. '.
3. Interactive Navigation:
   - Hyperlinked Table of Contents.
   - Standardized return link ('⬅ Return to Table of Contents') in Cell A1.
4. Auto-adjusted column widths and intelligent value parsing (integers, floats, percentages, alignments).
"""

from typing import Any, Tuple, Optional, Dict
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# Border definitions (APA 7th Edition compliant)
SIDE_THIN = Side(border_style="thin", color="000000")
BORDER_NONE = Border(left=None, right=None, top=None, bottom=None)
BORDER_TOP_BOTTOM = Border(top=SIDE_THIN, bottom=SIDE_THIN, left=None, right=None)
BORDER_BOTTOM_ONLY = Border(bottom=SIDE_THIN, left=None, right=None, top=None)
BORDER_TOP_ONLY = Border(top=SIDE_THIN, left=None, right=None, bottom=None)

# Standard Fonts (Calibri)
FONT_TABLE_ID = Font(name="Calibri", size=11, bold=True, color="000000")
FONT_TABLE_TITLE = Font(name="Calibri", size=11, italic=True, color="000000")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="000000")
FONT_DATA = Font(name="Calibri", size=11, bold=False, italic=False, color="000000")
FONT_NOTE = Font(name="Calibri", size=10, italic=False, color="333333")
FONT_NOTE_PREFIX = Font(name="Calibri", size=10, italic=True, color="333333")
FONT_LINK = Font(name="Calibri", size=11, color="0563C1", underline="single")

# Fills
FILL_NONE = PatternFill(fill_type=None)


def parse_cell_value(val_str: str) -> Tuple[Any, Optional[str], str]:
    """
    Parses a raw CSV string value into an appropriate Python type, Excel number format,
    and horizontal alignment.
    
    Returns:
        (parsed_value, number_format, horizontal_alignment)
    """
    s = str(val_str).strip()
    if not s:
        return "", None, "left"

    # Missing / Null representations or dashes
    if s in ("—", "-", "--", "N/A", "NA", "None", "NULL", "nan"):
        return s, None, "center"

    # Boolean values
    if s.upper() in ("TRUE", "FALSE", "YES", "NO"):
        return s, None, "center"

    # Formatted integer with commas e.g. '24,678,912'
    clean_int_s = s.replace(",", "")
    if clean_int_s.isdigit() and ("," in s or len(s) > 1):
        try:
            return int(clean_int_s), "#,##0", "right"
        except ValueError:
            pass

    # Pure signed integer
    if s.isdigit() or (s.startswith("-") and s[1:].isdigit()):
        try:
            return int(s), "#,##0", "right"
        except ValueError:
            pass

    # Floating point numbers (handling decimals)
    try:
        f = float(s)
        if "." in s:
            decimals = len(s.split(".")[1])
            fmt = "0." + "0" * min(decimals, 4)
        else:
            fmt = "0.00"
        return f, fmt, "right"
    except ValueError:
        pass

    # Percentages e.g. '+108.4%' or '0.0%'
    if s.endswith("%"):
        clean_num = s[:-1].replace("+", "").strip()
        try:
            float(clean_num)
            return s, None, "right"
        except ValueError:
            pass

    # Checkpoint codes / Time stamps (e.g. '05:00') / Airport codes (e.g. 'DFW', 'BOS')
    if len(s) <= 5 and (s.isupper() or ":" in s or s in ("BP01", "LAN", "SBP")):
        return s, None, "center"

    return s, None, "left"


def auto_fit_columns(ws, min_col: int = 1, max_col: Optional[int] = None,
                     min_row: int = 1, max_row: Optional[int] = None,
                     min_width: int = 12, max_width: int = 60, padding: int = 4):
    """
    Adjusts worksheet column widths based on maximum cell string lengths within the specified range.
    """
    max_c = max_col or ws.max_column
    max_r = max_row or ws.max_row
    
    for col_idx in range(min_col, max_c + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(min_row, max_r + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val_s = str(cell.value or "")
            if len(val_s) > max_len:
                max_len = len(val_s)
        ws.column_dimensions[col_letter].width = max(min(max_len + padding, max_width), min_width)


def setup_return_link(ws, row: int = 1, col: int = 1,
                      target: str = "#'Contents'!A1",
                      text: str = "⬅ Return to Table of Contents"):
    """
    Inserts a standardized APA 7th Edition hyperlink back to the Table of Contents.
    """
    c = ws.cell(row=row, column=col, value=text)
    c.font = FONT_LINK
    c.hyperlink = target
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20
