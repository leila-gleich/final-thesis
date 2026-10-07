"""
tests/test_figure_workbooks.py
------------------------------
Unit tests verifying the generation, integrity, and APA 7th Edition compliance
of the 4 consolidated Excel workbooks created from the figures/ subfolders.
"""

import unittest
from pathlib import Path
import openpyxl
from src.analysis.generate_figures_tables_excel import (
    FIGURES_DIR,
    SUBFOLDER_WORKBOOKS,
    generate_workbook_for_subfolder,
    build_all_figure_workbooks,
)


class TestFigureWorkbooks(unittest.TestCase):
    def test_build_all_figure_workbooks(self):
        """Verifies that all 4 figure workbooks build successfully with APA 7 rules."""
        generated = build_all_figure_workbooks()
        self.assertEqual(len(generated), 4)

        for sf_info in SUBFOLDER_WORKBOOKS:
            folder_name = sf_info["folder_name"]
            wb_filename = sf_info["wb_filename"]
            expected_tables = sf_info["tables"]
            target_dir = sf_info.get("target_dir", FIGURES_DIR / folder_name)
            
            wb_path = target_dir / wb_filename
            self.assertTrue(wb_path.exists(), f"Workbook {wb_path} does not exist.")

            wb = openpyxl.load_workbook(wb_path)
            
            # 1. Verify Contents sheet exists as first tab
            self.assertEqual(wb.sheetnames[0], "Contents")
            ws_contents = wb["Contents"]
            self.assertIn("TABLES", ws_contents["A1"].value.upper())
            
            # 2. Verify all expected table sheets exist
            for t_info in expected_tables:
                sname = t_info["short_name"]
                self.assertIn(sname, wb.sheetnames, f"Sheet {sname} missing in {wb_filename}")
                ws_table = wb[sname]

                # APA 7: No auto-filters or Table objects
                self.assertIsNone(ws_table.auto_filter.ref, f"AutoFilter found in {sname}")
                self.assertEqual(len(ws_table.tables), 0, f"Table object found in {sname}")

                # APA 7: Return link in row 1
                self.assertEqual(ws_table.cell(1, 1).value, "⬅ Return to Table of Contents")
                self.assertEqual(ws_table.cell(1, 1).hyperlink.target, "#'Contents'!A1")

                # APA 7: Table ID bold on row 2, Title italic on row 3
                self.assertTrue(ws_table.cell(2, 1).font.bold)
                self.assertTrue(ws_table.cell(3, 1).font.italic)

                # APA 7: Header row on row 5 has top and bottom thin black border, no vertical rules
                h_cell = ws_table.cell(5, 1)
                self.assertIsNotNone(h_cell.border.top)
                self.assertEqual(h_cell.border.top.style, "thin")
                self.assertIsNotNone(h_cell.border.bottom)
                self.assertEqual(h_cell.border.bottom.style, "thin")
                self.assertTrue(h_cell.border.left is None or h_cell.border.left.style is None)
                self.assertTrue(h_cell.border.right is None or h_cell.border.right.style is None)

                # APA 7: Note exists below the table
                max_r = ws_table.max_row
                note_val = str(ws_table.cell(max_r, 1).value or "")
                self.assertTrue(note_val.startswith("Note. "), f"Missing Note. in {sname}: {note_val}")
                self.assertTrue(ws_table.cell(max_r, 1).font.italic)


if __name__ == "__main__":
    unittest.main()
