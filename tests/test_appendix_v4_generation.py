#!/usr/bin/env python3
"""
tests/test_appendix_v4_generation.py
====================================
Unit tests verifying the integrity, completeness, and APA 7th Edition compliance
of Appendix v4 generation across both multi-tab Excel and Markdown targets.
"""

import os
import unittest
import openpyxl

from src.manuscripts.generate_appendix_v4 import (
    TABLES_CONFIG,
    create_appendix_excel_workbook,
    generate_appendix_v4_markdown
)


class TestAppendixV4Generation(unittest.TestCase):
    """Test suite for Appendix v4 compilation."""

    @classmethod
    def setUpClass(cls):
        cls.excel_path = "thesis_docs/manuscripts/Appendix/Appendix_Tables.xlsx"
        cls.results_excel_path = "results/appendix_tables.xlsx"
        cls.md_path = "thesis_docs/manuscripts/Appendix/Appendix v4.md"
        cls.root_md_path = "thesis_docs/manuscripts/Appendix v4.md"
        cls.appendix_md_path = "thesis_docs/manuscripts/Appendix/appendix.md"

    def test_excel_workbooks_exist_and_tabs_complete(self):
        """Verify that both Excel workbooks exist and contain exactly 30 tabs."""
        for path in [self.excel_path, self.results_excel_path]:
            self.assertTrue(os.path.exists(path), f"Excel workbook missing: {path}")
            wb = openpyxl.load_workbook(path, read_only=True)
            self.assertEqual(len(wb.sheetnames), 30, f"Expected 30 sheets in {path}, found {len(wb.sheetnames)}")
            self.assertIn("Contents", wb.sheetnames)
            self.assertIn("Table_A1_Assumptions", wb.sheetnames)
            self.assertIn("Table_B1_Paradigms", wb.sheetnames)
            self.assertIn("Table_B2_Equations", wb.sheetnames)
            self.assertIn("Table_C1_Queuing", wb.sheetnames)
            self.assertIn("Table_D1_Partitioning", wb.sheetnames)
            self.assertIn("Table_E1_Top25_Census", wb.sheetnames)
            self.assertIn("Table_F1_Data_Feeds", wb.sheetnames)
            self.assertIn("Table_G1_Statistical_Tests", wb.sheetnames)
            self.assertIn("Table_H1_Feature_Dictionary", wb.sheetnames)
            self.assertIn("Table_I1_Seasonal_Regimes", wb.sheetnames)
            self.assertIn("Table_J1_Evaluation_Targets", wb.sheetnames)
            self.assertIn("Table_K1_Values_vs_Volatility", wb.sheetnames)
            self.assertIn("Table_L1_Disruption_Resilience", wb.sheetnames)
            self.assertIn("Table_M1_Routine_Accuracy", wb.sheetnames)
            self.assertIn("Table_N1_Archival_Hierarchy", wb.sheetnames)

    def test_markdown_files_exist_and_cover_all_appendices(self):
        """Verify that all Markdown targets exist and cover Appendices A through N."""
        for path in [self.md_path, self.root_md_path, self.appendix_md_path]:
            self.assertTrue(os.path.exists(path), f"Markdown file missing: {path}")
            with open(path, "r", encoding="utf-8") as f:
                content = f.read()
            self.assertIn("# Appendix A: Delimitations", content)
            self.assertIn("# Appendix B: Comparative Modeling Paradigms", content)
            self.assertIn("# Appendix C: Queuing Theory Derivations", content)
            self.assertIn("# Appendix D: Temporal Scope", content)
            self.assertIn("# Appendix E: Four-Tier Purposive Filtering Pipeline", content)
            self.assertIn("# Appendix F: Multi-Source Aviation Data Feeds", content)
            self.assertIn("# Appendix G: Statistical Significance Testing", content)
            self.assertIn("# Appendix H: Feature Engineering Pipeline", content)
            self.assertIn("# Appendix I: Seasonal Volatility Regimes", content)
            self.assertIn("# Appendix J: Candidate Predictive Models", content)
            self.assertIn("# Appendix K: The Values versus Volatility Operational Coupling", content)
            self.assertIn("# Appendix L: Operational Disruption Dynamics", content)
            self.assertIn("# Appendix M: Dual-Track Operational Decision Playbook", content)
            self.assertIn("# Appendix N: Archival Data Storage Infrastructure", content)

    def test_policy_compliance_in_markdown(self):
        """Verify strict adherence to thesis constitution policies in markdown."""
        with open(self.md_path, "r", encoding="utf-8") as f:
            content = f.read()

        # Policy 2.1: Throughput Volatility targeting
        self.assertIn("volatility", content.lower())
        self.assertIn("\\sigma_{\\text{TSA", content)
        self.assertIn("CV_{\\text{TSA", content)

        # Policy 3.1 & 3.2: 3 models + baseline control, zero internal code tags
        self.assertIn("Baseline Control", content)
        self.assertIn("Model 1", content)
        self.assertIn("Model 2", content)
        self.assertIn("Model 3", content)
        self.assertNotIn("M1*", content)
        self.assertNotIn("M0", content)
        self.assertNotIn("M5", content)

        # Policy 5.3: Authentic aviation terminology
        self.assertIn("Nominal On-Time Baseline", content)
        self.assertIn("Routine Daily Operations", content)
        self.assertIn("Irregular Operations", content)
        self.assertNotIn("quiescent", content.lower())
        self.assertNotIn("sterile control", content.lower())

    def test_table_references_integrity(self):
        """Verify that every configured table is cited with its Excel tab in markdown."""
        with open(self.md_path, "r", encoding="utf-8") as f:
            content = f.read()
        for cfg in TABLES_CONFIG:
            self.assertIn(f"### {cfg['table_id']}", content)
            self.assertIn(cfg["sheet_name"], content)


if __name__ == "__main__":
    unittest.main()
