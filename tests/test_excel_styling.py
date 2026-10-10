"""
tests/test_excel_styling.py
---------------------------
Unit tests for centralized Excel APA 7th Edition styling and numeric parsing utilities.
"""

import os
import sys
import unittest
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
from utils.excel_styling import (
    parse_cell_value,
    apply_apa_table_borders,
    style_table_range,
    autofit_column_widths,
    APA_FONT_NAME
)

class TestExcelStyling(unittest.TestCase):
    def test_parse_cell_value_integers(self):
        val, fmt, align = parse_cell_value("100")
        self.assertEqual(val, 100)
        self.assertEqual(fmt, "#,##0")
        self.assertEqual(align, "right")

        val, fmt, align = parse_cell_value("1,234")
        self.assertEqual(val, 1234)
        self.assertEqual(fmt, "#,##0")
        self.assertEqual(align, "right")

    def test_parse_cell_value_floats(self):
        val, fmt, align = parse_cell_value("12.34")
        self.assertAlmostEqual(val, 12.34)
        self.assertEqual(align, "right")

        val, fmt, align = parse_cell_value("1,234.56")
        self.assertAlmostEqual(val, 1234.56)
        self.assertEqual(align, "right")

        val, fmt, align = parse_cell_value(-8.5)
        self.assertAlmostEqual(val, -8.5)
        self.assertEqual(align, "right")

    def test_parse_cell_value_percentages(self):
        val, fmt, align = parse_cell_value("45.2%")
        self.assertEqual(val, "45.2%")
        self.assertEqual(align, "right")

    def test_parse_cell_value_strings_and_empty(self):
        val, fmt, align = parse_cell_value("Boston")
        self.assertEqual(val, "Boston")
        self.assertIsNone(fmt)
        self.assertEqual(align, "left")

        val, fmt, align = parse_cell_value("")
        self.assertEqual(val, "")
        self.assertIsNone(fmt)
        self.assertEqual(align, "left")

        val, fmt, align = parse_cell_value("—")
        self.assertEqual(val, "—")
        self.assertEqual(align, "center")

    def test_apa_table_borders_and_styling(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TestSheet"

        # Populate a sample 3x3 table
        data = [
            ["Model", "RMSE", "MASE"],
            ["Model 1", 245.5, 0.945],
            ["Model 3", 222.1, 0.662]
        ]
        for r_idx, row in enumerate(data, 1):
            for c_idx, val in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)

        # Apply APA styling
        apply_apa_table_borders(ws, header_row_idx=1, data_start_row=2, data_end_row=3, start_col=1, end_col=3)
        style_table_range(ws, start_row=1, end_row=3, start_col=1, end_col=3, font_size=10)
        autofit_column_widths(ws, min_col=1, max_col=3, min_width=10, max_width=30)

        # Verify header top and bottom border
        header_cell = ws.cell(row=1, column=1)
        self.assertEqual(header_cell.border.top.style, "thin")
        self.assertEqual(header_cell.border.bottom.style, "thin")
        # Verify no vertical border on header
        self.assertTrue(header_cell.border.left is None or header_cell.border.left.style is None)
        self.assertTrue(header_cell.border.right is None or header_cell.border.right.style is None)

        # Verify last row bottom border
        last_cell = ws.cell(row=3, column=2)
        self.assertEqual(last_cell.border.bottom.style, "thin")

        # Verify typography
        self.assertEqual(header_cell.font.name, APA_FONT_NAME)
        self.assertEqual(last_cell.font.name, APA_FONT_NAME)

        # Verify column widths adjusted
        self.assertGreaterEqual(ws.column_dimensions["A"].width, 10)
        self.assertLessEqual(ws.column_dimensions["A"].width, 30)

if __name__ == "__main__":
    unittest.main()
